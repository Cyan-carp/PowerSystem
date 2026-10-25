import hashlib
import importlib.util
import json
from pathlib import Path
import httpx
import pytest
from rag_agent.knowledge import Knowledge
from test_engine import config

spec = importlib.util.spec_from_file_location("import_knowledge", Path(__file__).parents[1] / "scripts/knowledge.py")
ingest = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ingest)

audit_spec = importlib.util.spec_from_file_location("audit_knowledge", Path(__file__).parents[1] / "scripts/audit-v4-knowledge.py")
audit = importlib.util.module_from_spec(audit_spec)
audit_spec.loader.exec_module(audit)


class Tokenizer:
    def encode(self, text, **kwargs): return list(text)
    def decode(self, tokens): return "".join(tokens)


def test_chunks_bound_tokens_and_preserve_tail():
    text = "甲" * 1000 + "最后一段"
    rows = list(ingest.chunks([("标题", text, 7)], Tokenizer()))
    assert len(rows) > 2 and all(len(row[1]) <= 400 for row in rows)
    assert rows[-1][1].endswith("最后一段\n") and all(row[2] == 7 for row in rows)
    assert rows[0][1][-60:] in rows[1][1]


@pytest.mark.asyncio
async def test_historical_source_reads_retained_collection():
    cfg = config()
    def handler(request):
        path = request.url.path
        if path == "/collections":
            return httpx.Response(200, json={"result":{"collections":[{"name":"unrelated_private"},{"name":"powersystem_rag_old_123"}]}})
        if "powersystem_rag_old_123" in path:
            return httpx.Response(200, json={"result":{"payload":{"version":"old","content":"历史章节"}}})
        if "unrelated_private" in path: pytest.fail("must not search unrelated collections")
        return httpx.Response(404)
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        assert (await Knowledge(cfg, client).source("K" + "a" * 20))["version"] == "old"


def test_extract_document_types():
    # Keep generated files in the ignored test assets directory.
    folder = Path("artifacts/stage8-agent-rag/test-documents")
    folder.mkdir(parents=True, exist_ok=True)
    from docx import Document
    document = Document(); document.add_heading("设备资料", 1); document.add_paragraph("风扇用于散热")
    table = document.add_table(rows=1, cols=2); table.cell(0,0).text = "型号"; table.cell(0,1).text = "样例"
    document.save(folder / "device.docx")
    assert "风扇" in str(ingest.extract(folder / "device.docx")) and "样例" in str(ingest.extract(folder / "device.docx"))
    (folder / "sample.md").write_text("# 标题\n\n测试段落", encoding="utf-8")
    (folder / "sample.txt").write_text("测试文本", encoding="utf-8")
    assert ingest.extract(folder / "sample.md")[0][0] == "标题"
    assert "测试文本" in str(ingest.extract(folder / "sample.txt"))
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter(); page = writer.add_blank_page(width=200, height=200)
    font = DictionaryObject({NameObject("/Type"):NameObject("/Font"), NameObject("/Subtype"):NameObject("/Type1"),NameObject("/BaseFont"):NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):DictionaryObject({NameObject("/F1"):writer._add_object(font)})})
    stream = DecodedStreamObject(); stream.set_data(b"BT /F1 12 Tf 10 100 Td (Cooling manual) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_blank_page(width=200,height=200)
    writer.write(folder / "manual.pdf")
    sections = ingest.extract(folder / "manual.pdf")
    assert "Cooling manual" in sections[0][1] and sections[0][2] == 1
    assert sections[1][1] == "" and not list(ingest.chunks([sections[1]], Tokenizer()))


@pytest.mark.asyncio
async def test_search_keeps_source_version_hash_and_collection():
    cfg = config()
    class Embedding:
        def encode(self, texts, query=False): return [[0.1, 0.2]]
    def handler(request):
        if request.url.path == '/aliases':
            return httpx.Response(200,json={'result':{'aliases':[{'alias_name':cfg.collection,'collection_name':'powersystem_rag_v1'}]}})
        assert request.url.path.endswith("/points/query")
        assert '/powersystem_rag_v1/' in request.url.path
        return httpx.Response(200, json={"result": {"points": [{"score": 0.9, "payload": {
            "id": "K" + "a" * 20, "document": "手册.md", "heading": "章节",
            "version": "v1", "source_hash": "h" * 64}}]}})
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        rows = await Knowledge(cfg, client, Embedding()).search("告警")
    assert rows[0]["collection"] == 'powersystem_rag_v1'
    assert (rows[0]["document"], rows[0]["version"], rows[0]["source_hash"]) == ("手册.md", "v1", "h" * 64)


@pytest.mark.asyncio
async def test_historical_source_reports_collection_name():
    cfg = config()
    def handler(request):
        path = request.url.path
        if path == "/collections":
            return httpx.Response(200, json={"result": {"collections": [{"name": "powersystem_rag_old_123"}]}})
        if "powersystem_rag_old_123" in path:
            return httpx.Response(200, json={"result": {"payload": {"version": "old"}}})
        return httpx.Response(404)
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        row = await Knowledge(cfg, client).source("K" + "a" * 20)
    assert row["collection"] == "powersystem_rag_old_123"


def test_manifest_digest_is_repeatable_and_content_sensitive():
    entries = [{"path": "a.md", "sha256": "x"}]
    assert ingest.manifest_sha256(entries) == ingest.manifest_sha256(list(entries))
    assert ingest.manifest_sha256(entries) != ingest.manifest_sha256([{"path": "a.md", "sha256": "y"}])


def test_document_hash_survives_deployment_line_ending_normalization(tmp_path):
    path=tmp_path/'a.md'
    path.write_bytes('标题\r\n正文\r\n'.encode('utf-8'))
    before=ingest.document_sha256(path)
    path.write_bytes('标题\n正文\n'.encode('utf-8'))
    assert ingest.document_sha256(path)==before


def test_collect_is_repeatable_and_pins_chunk_index(tmp_path):
    (tmp_path / "a.md").write_text("# 标题\n\n第一段\n\n第二段", encoding="utf-8")
    entry = {"path": "a.md", "title": "A", "sha256": ingest.document_sha256(tmp_path / "a.md")}
    first = ingest.collect([entry], Tokenizer(), root=tmp_path)
    second = ingest.collect([entry], Tokenizer(), root=tmp_path)
    assert not first[4] and first[2] == []
    assert [row["id"] for row in first[0]] == [row["id"] for row in second[0]]
    assert first[3] == second[3]
    assert first[3][0]["content_sha256"] == hashlib.sha256(first[0][0]["content"].encode()).hexdigest()


def test_collect_blocks_manifest_hash_drift(tmp_path):
    (tmp_path / "a.md").write_text("# 标题\n\n正文", encoding="utf-8")
    entry = {"path": "a.md", "title": "A", "sha256": "0" * 64}
    rows, files, _, _, blockers = ingest.collect([entry], Tokenizer(), root=tmp_path)
    assert rows == [] and files == []
    assert blockers[0]["code"] == "document_hash_drift"
    rows, files, _, _, blockers = ingest.collect([entry], Tokenizer(), root=tmp_path, verify_manifest=False)
    assert rows and blockers[0]["code"] == "document_hash_drift"


def test_collect_reports_missing_document(tmp_path):
    _, _, _, _, blockers = ingest.collect([{"path": "gone.md"}], Tokenizer(), root=tmp_path)
    assert blockers == [{"code": "document_missing", "document": "gone.md"}]


def test_audit_documents_uses_content_hash_not_mtime(tmp_path):
    path = tmp_path / "a.md"
    path.write_text("# 标题\n\n正文", encoding="utf-8")
    good = {"path": "a.md", "title": "A", "sha256": ingest.document_sha256(path)}
    documents, blockers, warnings = audit.audit_documents([good], ingest, Tokenizer(), tmp_path)
    assert not blockers and documents[0]["hash_status"] == "match"
    assert documents[0]["chunks"] == 1 and warnings == []
    path.write_text("# 标题\n\n改过的正文", encoding="utf-8")
    documents, blockers, _ = audit.audit_documents([good], ingest, Tokenizer(), tmp_path)
    assert documents[0]["hash_status"] == "drift"
    assert blockers[0]["code"] == "document_hash_drift"


def test_audit_documents_flags_missing_manifest_hash(tmp_path):
    (tmp_path / "a.md").write_text("# 标题\n\n正文", encoding="utf-8")
    documents, blockers, warnings = audit.audit_documents([{"path": "a.md"}], ingest, Tokenizer(), tmp_path)
    assert not blockers and documents[0]["hash_status"] == "manifest_hash_missing"
    assert warnings[0]["code"] == "manifest_hash_missing"


def collection_client(points, aliases=None, count=None):
    aliases = aliases if aliases is not None else [{"alias_name": "powersystem_rag_active",
                                                    "collection_name": "powersystem_rag_v1_x"}]
    def handler(request):
        path = request.url.path
        if path == "/aliases":
            return httpx.Response(200, json={"result": {"aliases": aliases}})
        if path.endswith("/points/count"):
            return httpx.Response(200, json={"result": {"count": count if count is not None else len(points)}})
        if path.endswith("/points/scroll"):
            return httpx.Response(200, json={"result": {"points": points, "next_page_offset": None}})
        return httpx.Response(404)
    return httpx.Client(transport=httpx.MockTransport(handler), base_url="http://qdrant")


def test_audit_collection_matches_active_payload():
    points = [{"payload": {"document": "A", "source_hash": "h" * 64, "version": "v1"}},
              {"payload": {"document": "A", "source_hash": "h" * 64, "version": "v1"}}]
    with collection_client(points) as client:
        result, blockers = audit.audit_collection(client, "powersystem_rag_active",
            [{"document": "A", "sha256": "h" * 64, "chunks": 2}], "v1")
    assert result["resolved"] == "powersystem_rag_v1_x" and result["points"] == 2
    assert blockers == [] and result["diff"]["missing_documents"] == []


def test_audit_collection_reports_structured_diff():
    points = [{"payload": {"document": "B", "source_hash": "x", "version": "old"}}]
    with collection_client(points) as client:
        result, blockers = audit.audit_collection(client, "powersystem_rag_active",
            [{"document": "A", "sha256": "h" * 64, "chunks": 1}], "v1")
    codes = {row["code"] for row in blockers}
    assert {"active_document_missing", "active_document_unexpected"} <= codes
    assert result["diff"]["missing_documents"] == ["A"]
    assert result["diff"]["extra_documents"] == ["B"]


def test_audit_collection_reports_missing_alias():
    with collection_client([], aliases=[]) as client:
        result, blockers = audit.audit_collection(client, "powersystem_rag_active", [], None)
    assert blockers[0]["code"] == "collection_alias_missing" and result["exists"] is False


def test_safe_origin_strips_credentials_and_path():
    assert audit.safe_origin("http://user:pass@127.0.0.1:6333/collections?x=1") == "http://127.0.0.1:6333"
    assert audit.safe_origin("not-a-url") == ""


@pytest.mark.asyncio
async def test_pinned_old_import_does_not_resolve_same_chunk_id_from_new_alias():
    cfg=config();seen=[]
    def handler(request):
        seen.append(request.url.path)
        assert '/powersystem_rag_old_import/' in request.url.path
        return httpx.Response(200,json={'result':{'payload':{'id':'K'+'a'*20,'version':'a'*64,'content':'same bytes'}}})
    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        k=Knowledge(cfg,client)
        old=await k.source('K'+'a'*20,'powersystem_rag_old_import','a'*64)
        assert old['version']=='a'*64 and old['collection']=='powersystem_rag_old_import'
        assert await k.source('K'+'a'*20,'powersystem_rag_old_import','b'*64) is None
        for bad in ('../private','unrelated_private','powersystem_rag_active'):
            with pytest.raises(ValueError):await k.source('K'+'a'*20,bad,'a'*64)
    assert len(seen)==2
